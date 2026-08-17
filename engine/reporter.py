"""
Report Generation Engine
Converts analysis results into a structured research report.
Supports plain text markdown (for Streamlit display) and DOCX download.
"""

import io
from datetime import datetime
import pandas as pd


def generate_report_markdown(
    dataset_name: str,
    profile: dict,
    cleaning_log: list,
    analyses: list[dict],
    key_findings: list[str],
) -> str:
    """
    Generate a full research report as a Markdown string.
    Every claim is traced to its source: data, calculation, result, interpretation.
    """
    lines = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    lines.append("# Research Data Analysis Report")
    lines.append(f"**Generated:** {now}  ")
    lines.append(f"**Dataset:** {dataset_name}")
    lines.append("")
    lines.append("---")

    # ── 1. Dataset Description ───────────────────────────────────────────────
    lines.append("")
    lines.append("## 1. Dataset Description")
    lines.append("")
    lines.append(f"- **Total rows:** {profile.get('n_rows', 'N/A'):,}")
    lines.append(f"- **Total columns:** {profile.get('n_cols', 'N/A')}")
    n_cols = profile.get("numerical_cols", [])
    c_cols = profile.get("categorical_cols", [])
    d_cols = profile.get("datetime_cols", [])
    if n_cols:
        lines.append(f"- **Numerical variables ({len(n_cols)}):** {', '.join(n_cols)}")
    if c_cols:
        lines.append(f"- **Categorical variables ({len(c_cols)}):** {', '.join(c_cols)}")
    if d_cols:
        lines.append(f"- **Date/Time variables ({len(d_cols)}):** {', '.join(d_cols)}")
    lines.append(f"- **Missing values:** {profile.get('total_missing', 0):,} "
                 f"({profile.get('total_missing_pct', 0):.2f}% of all cells)")
    lines.append(f"- **Duplicate rows detected:** {profile.get('duplicate_rows', 0)}")
    lines.append("")

    # ── 2. Data Cleaning Report ───────────────────────────────────────────────
    lines.append("## 2. Data Cleaning Report")
    lines.append("")
    if not cleaning_log:
        lines.append("*No data cleaning operations were performed on this dataset.*")
    else:
        lines.append(f"**{len(cleaning_log)} cleaning operation(s) performed.**")
        lines.append("")
        lines.append("| # | Timestamp | Operation | Column | Method | Values Changed | Detail |")
        lines.append("|---|-----------|-----------|--------|--------|----------------|--------|")
        for i, entry in enumerate(cleaning_log, 1):
            lines.append(
                f"| {i} | {entry.get('timestamp','')} | {entry.get('operation','')} | "
                f"`{entry.get('column','')}` | {entry.get('method','')} | "
                f"{entry.get('values_changed', 0)} | {entry.get('detail', '')} |"
            )
        lines.append("")
        lines.append("> **Note:** The original uploaded dataset was preserved unchanged. "
                     "All cleaning operations were applied to a separate working copy.")
    lines.append("")

    # ── 3. Statistical Analyses ───────────────────────────────────────────────
    lines.append("## 3. Statistical Analysis Results")
    lines.append("")
    if not analyses:
        lines.append("*No statistical analyses have been performed yet.*")
    else:
        for i, result in enumerate(analyses, 1):
            test = result.get("test", f"Analysis {i}")
            lines.append(f"### 3.{i} {test}")
            lines.append("")

            # Evidence chain
            lines.append("**Evidence Chain:**")
            lines.append(f"- Raw data used: columns `{result.get('col1', result.get('num_col', result.get('dependent', '')))}` "
                         f"and `{result.get('col2', result.get('group_col', ''))}`")
            lines.append(f"- Sample size (n): {result.get('n', 'N/A')}")

            # Core statistics
            if "r" in result:
                lines.append(f"- Correlation coefficient (r): **{result['r']}**")
            if "chi2" in result:
                lines.append(f"- Chi-square statistic: **{result['chi2']}** (df = {result.get('df', 'N/A')})")
                lines.append(f"- Cramér's V (effect size): **{result.get('cramers_v', 'N/A')}**")
            if "t_statistic" in result:
                lines.append(f"- t-statistic: **{result['t_statistic']}**")
                lines.append(f"- Cohen's d (effect size): **{result.get('cohens_d', 'N/A')}**")
            if "f_statistic" in result:
                lines.append(f"- F-statistic: **{result['f_statistic']}**")
                lines.append(f"- η² (effect size): **{result.get('eta_squared', 'N/A')}**")
            if "r_squared" in result:
                lines.append(f"- R²: **{result['r_squared']}** (Adjusted R²: {result.get('adj_r_squared', 'N/A')})")

            if "p_value" in result:
                p = result["p_value"]
                sig = "**Statistically significant** (p < 0.05)" if p < 0.05 else "Not statistically significant (p ≥ 0.05)"
                lines.append(f"- p-value: **{p}** — {sig}")

            lines.append("")
            lines.append(f"**Interpretation:** {result.get('interpretation', '')}")
            lines.append("")
            lines.append(f"**Assumptions:** {result.get('assumption_note', '')}")
            if result.get("warning"):
                lines.append(f"⚠️ **Warning:** {result['warning']}")
            lines.append("")

    # ── 4. Key Findings ───────────────────────────────────────────────────────
    lines.append("## 4. Key Findings")
    lines.append("")
    if not key_findings:
        lines.append("*Key findings will appear here as analyses are completed.*")
    else:
        for finding in key_findings:
            lines.append(f"- {finding}")
    lines.append("")

    # ── 5. Limitations ────────────────────────────────────────────────────────
    lines.append("## 5. Limitations and Methodological Considerations")
    lines.append("")
    lines.append("The following limitations should be considered when interpreting the results:")
    lines.append("")
    n_rows = profile.get("n_rows", 0)
    if n_rows < 30:
        lines.append(f"- ⚠️ **Small sample size** (n = {n_rows}): Results should be interpreted with caution. "
                     "Statistical tests have lower power and results may not be generalisable.")
    if profile.get("total_missing", 0) > 0:
        lines.append(f"- ⚠️ **Missing data** ({profile.get('total_missing', 0)} values): Imputation methods "
                     "introduce assumptions. The pattern of missing data (MCAR, MAR, MNAR) may affect results.")
    if cleaning_log:
        lines.append("- ⚠️ **Data cleaning was performed**: Cleaning operations (especially imputation and outlier handling) "
                     "modify the dataset and can influence statistical results.")
    lines.append("- Statistical significance (p < 0.05) does not imply practical significance or causation.")
    lines.append("- Results are specific to this dataset and should not be over-generalised.")
    lines.append("")

    # ── 6. Research Transparency ──────────────────────────────────────────────
    lines.append("## 6. Research Transparency")
    lines.append("")
    lines.append("This report was generated from verified computational results using the following pipeline:")
    lines.append("")
    lines.append("```")
    lines.append("Original Data Upload")
    lines.append("  → Data Profiling (pandas)")
    lines.append("  → Data Cleaning (documented in Section 2)")
    lines.append("  → Statistical Engine (scipy, statsmodels)")
    lines.append("  → AI Interpretation Layer (explanation only — does not alter computed values)")
    lines.append("  → Report Generation")
    lines.append("```")
    lines.append("")
    lines.append("*All statistical values in this report are computed by established Python statistical libraries. "
                 "The AI layer provides explanations and interpretations but does not calculate or modify statistical results.*")

    return "\n".join(lines)


def generate_report_docx(
    dataset_name: str,
    profile: dict,
    cleaning_log: list,
    analyses: list[dict],
    key_findings: list[str],
) -> bytes:
    """
    Generate a DOCX report. Returns the raw bytes.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        now = datetime.now().strftime("%Y-%m-%d %H:%M")

        # Title
        title = doc.add_heading("Research Data Analysis Report", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Generated: {now}  |  Dataset: {dataset_name}")

        doc.add_heading("1. Dataset Description", level=1)
        p = doc.add_paragraph()
        p.add_run(f"Rows: {profile.get('n_rows', 'N/A'):,}  |  "
                  f"Columns: {profile.get('n_cols', 'N/A')}  |  "
                  f"Missing: {profile.get('total_missing', 0):,} ({profile.get('total_missing_pct', 0):.2f}%)  |  "
                  f"Duplicates: {profile.get('duplicate_rows', 0)}")

        doc.add_heading("2. Data Cleaning", level=1)
        if not cleaning_log:
            doc.add_paragraph("No cleaning operations performed.")
        else:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Timestamp", "Operation", "Column", "Method", "Detail"]):
                hdr[i].text = h
            for entry in cleaning_log:
                row = table.add_row().cells
                row[0].text = str(entry.get("timestamp", ""))
                row[1].text = str(entry.get("operation", ""))
                row[2].text = str(entry.get("column", ""))
                row[3].text = str(entry.get("method", ""))
                row[4].text = str(entry.get("detail", ""))

        doc.add_heading("3. Statistical Results", level=1)
        for result in analyses:
            doc.add_heading(result.get("test", "Analysis"), level=2)
            if "r" in result:
                doc.add_paragraph(f"r = {result['r']}   p = {result.get('p_value', 'N/A')}   n = {result.get('n', 'N/A')}")
            if "t_statistic" in result:
                doc.add_paragraph(f"t = {result['t_statistic']}   p = {result.get('p_value', 'N/A')}")
            if "f_statistic" in result:
                doc.add_paragraph(f"F = {result['f_statistic']}   p = {result.get('p_value', 'N/A')}")
            if "chi2" in result:
                doc.add_paragraph(f"χ² = {result['chi2']}   df = {result.get('df', 'N/A')}   p = {result.get('p_value', 'N/A')}")
            doc.add_paragraph(f"Interpretation: {result.get('interpretation', '')}")
            doc.add_paragraph(f"Assumptions: {result.get('assumption_note', '')}")

        doc.add_heading("4. Key Findings", level=1)
        for f in key_findings:
            doc.add_paragraph(f, style="List Bullet")

        doc.add_heading("5. Limitations", level=1)
        doc.add_paragraph("Statistical significance does not imply causation or practical significance.")
        if profile.get("n_rows", 0) < 30:
            doc.add_paragraph(f"Small sample size (n = {profile.get('n_rows')}): Interpret with caution.")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    except ImportError:
        return b""
